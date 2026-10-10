"""BOOTSTRAP ONLY - turn workbook v3 into first-draft YAML definitions (Track B2).

    python tools/import_workbook.py --workbook <path to Cantier_Tyre_Datapoints_Ontology_Medallion_Catalogue_v3.xlsx>

Run ONCE to create the drafts. After that, the YAML files in Git are the source of truth and are edited by hand
through pull requests - re-running this script would overwrite those edits (it refuses unless --force is given).

What it writes (sheet -> folder):
  09 + 10            -> ontology/manufacturing.yaml           (LinkML, B2.1)
  03, 04A, 05, 06    -> glossary/measurement_types.yaml       (B2.2)
  entity rows + 09   -> glossary/entities.yaml                (B2.2)
  (fixed lists)      -> glossary/enums.yaml                   (B2.2)
  13 (other columns) -> glossary/columns.yaml                 (B2.4 needs a term for every business column)
  02 + 13            -> contracts/*.yaml                      (B2.3)
  12                 -> tables/catalogue.yaml                 (B2.4, every planned table)
  13                 -> tables/<layer>/<table>.yaml           (B2.4, tables that have columns)
  07 + 08            -> mappings/iiot/templates/*.yaml        (B2.5) + tests/reference/tag_instances_v3.csv
  04 Part B          -> mappings/mes_catalogue.yaml           (B2.6 input - mappings come with the MES schema export)
  07                 -> assets/site01.yaml                    (B2.7)
  table specs        -> dq/<table>.yaml                       (B2.8)
  04b                -> kpi/<area>.yaml                       (B2.9)
"""
import argparse
import csv
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parent.parent
VERSION = "0.1.0"
MODEL_VERSION = "manufacturing-v0.1"
MAPPING_VERSION = "iiot-map-2026.10.10"
HEADER = "# Draft generated from workbook v3 by tools/import_workbook.py - edit by pull request from now on.\n"


# ---------------------------------------------------------------- yaml output (keeps key order, readable)
class Dumper(yaml.SafeDumper):
    pass


Dumper.add_representer(OrderedDict, lambda d, data: d.represent_dict(data.items()))
Dumper.add_representer(dict, lambda d, data: d.represent_dict(data.items()))


def write_yaml(rel, data, comment=""):
    path = ROOT / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    text = yaml.dump(data, Dumper=Dumper, sort_keys=False, allow_unicode=True, width=140, default_flow_style=False)
    path.write_text(HEADER + comment + text, encoding="utf-8")
    WRITTEN.append(rel)


WRITTEN = []


# ---------------------------------------------------------------- small helpers
def snake(s):
    s = re.sub(r"[()\[\]]", " ", str(s))
    s = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", s)
    s = re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_").lower()
    return re.sub(r"_+", "_", s)


def human(tag):
    return snake(tag).replace("_", " ").capitalize()


def camel(s):
    return "".join(p[:1].upper() + p[1:] for p in re.split(r"[^A-Za-z0-9]+", s) if p)


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def rows(wb, sheet, header_contains):
    """Rows of a sheet as dicts, starting after the header row (first row whose first cell == header_contains)."""
    out, hdr = [], None
    for r in wb[sheet].iter_rows(values_only=True):
        if hdr is None:
            if r and r[0] == header_contains:
                hdr = [clean(h) for h in r]
            continue
        out.append({h: clean(v) for h, v in zip(hdr, r) if h})
    return out


DTYPE = {"double": "double", "float": "double", "int32": "int", "int16": "int", "int64": "bigint", "integer": "int",
         "boolean": "boolean", "string": "string", "datetime": "timestamp", "record": "string",
         "int": "int", "bigint": "bigint", "timestamp": "timestamp", "date": "date", "array<string>": "array<string>"}


def dtype(v):
    return DTYPE.get(str(v).strip().lower(), "string") if v else None


def physical_range(unit, data_type):
    """Only bounds that physics guarantees. Plant limits come from silver.mdm_spec_limit (decision 0.4)."""
    if data_type not in ("double", "int", "bigint") or not unit:
        return None
    if unit.startswith("%"):
        return {"min": 0, "max": 100}
    if unit in ("°C", "°F", "bar", "OPC UA StatusCode"):
        return None
    return {"min": 0}


def table_refs(text, layer):
    return sorted(set(re.findall(rf"\b{layer}\.[a-z0-9_]+", text or "")))


OWNER = {"IIoT": "IIoT / process engineering", "MES": "MES team", "inline quality": "Quality", "CMMS": "Maintenance",
         "PdM": "Data science", "Governance": "Platform governance"}


def owner_for(domain):
    for k, v in OWNER.items():
        if domain and k.lower() in domain.lower():
            return v
    return "Data platform team"


# ================================================================ B2.1 ontology (LinkML)
CLASS_NAMES = {  # sheet 09 "class" -> LinkML class name
    "Plant / Site": "Plant", "Process Area": "ProcessArea", "Production Line": "ProductionLine",
    "Asset (Machine)": "Asset", "Asset Type / Template": "AssetType", "Component": "Component",
    "Signal / Tag": "Signal", "Measurement Type": "MeasurementType", "Raw Material / Material Lot": "MaterialLot",
    "Compound Batch": "CompoundBatch", "Component (semi-finished)": "SemiFinishedComponent", "Green Tyre": "GreenTyre",
    "Cured Tyre": "CuredTyre", "Product (SKU)": "Product", "Recipe": "Recipe", "Spec Limit": "SpecLimit",
    "Production Order": "ProductionOrder", "Production Cycle": "ProductionCycle",
    "Machine State Event": "MachineStateEvent", "Alarm": "Alarm", "Downtime Event": "DowntimeEvent",
    "Inspection Result": "InspectionResult", "Lab Test": "LabTest", "Defect / NCR": "Defect",
    "Maintenance Work Order": "MaintenanceWorkOrder", "Failure Mode": "FailureMode", "Prediction": "Prediction",
    "Meter": "Meter", "Utility Supply": "UtilitySupply", "Shift": "Shift", "Approval / Audit": "Approval",
    "Lot (LotTraveller)": "Lot", "Lot Action": "LotAction", "Operation / Routing": "Operation",
    "Data Collection Parameter": "DataCollectionParameter", "Tool": "Tool", "Equipment Problem": "EquipmentProblem",
    "Inventory Location": "InventoryLocation", "Qualification": "Qualification", "Checklist Event": "ChecklistEvent",
}
# relation end names in sheet 10 -> classes (full-string matches first, they carry context)
REL_ENDS = {
    "Component (tread/ply)": ["SemiFinishedComponent"],
    "Compound Batch / Component / Green Tyre / Cured Tyre": ["CompoundBatch", "SemiFinishedComponent", "GreenTyre", "CuredTyre"],
    "Asset / Component": ["Asset", "Component"], "Raw Material Lot": ["MaterialLot"], "Signal reading": ["Signal"],
    "Signal readings": ["Signal"], "Signal readings (features)": ["Signal"], "Failure Mode on Component": ["FailureMode"],
    "Measurement": ["DataCollectionParameter"], "NCR": ["Defect"], "Production Order (JobOrder)": ["ProductionOrder"],
}
SEMI_FINISHED_CONTEXT = {("Component", "consumedBy", "Green Tyre")}  # here "Component" = tread / ply, not a machine part


def build_ontology(wb):
    by_label = {}
    for r in rows(wb, "09_Ontology_Classes", "class"):
        if r.get("class"):
            by_label[r["class"]] = r
    classes = OrderedDict()
    for label, r in by_label.items():
        name = CLASS_NAMES[label]
        ann = OrderedDict([("label", label), ("standard_alignment", r.get("standard alignment")),
                           ("id_prefix", r.get("ID prefix")), ("golden_source", r.get("golden source"))])
        classes[name] = OrderedDict([("description", r.get("definition / members in this plant")),
                                     ("annotations", OrderedDict((k, v) for k, v in ann.items() if v)),
                                     ("attributes", OrderedDict([("id", {"identifier": True, "description": f"Canonical id ({r.get('ID prefix') or 'n/a'}-...)"})]))])
    lookup = {k.lower(): v for k, v in {**{c: c for c in classes}, **{l: n for l, n in CLASS_NAMES.items()},
                                         "Plant": "Plant", "Asset Type": "AssetType", "Lot": "Lot"}.items()}

    def resolve(text, ctx=None):
        if ctx in SEMI_FINISHED_CONTEXT and text == "Component":
            return ["SemiFinishedComponent"]
        if text in REL_ENDS:
            return REL_ENDS[text]
        base = re.sub(r"\(.*?\)", "", text).strip()
        hit = lookup.get(base.lower()) or lookup.get(camel(base).lower())
        if not hit:
            raise SystemExit(f"ontology: cannot resolve relation end '{text}'")
        return [hit]

    n_rel = 0
    for r in rows(wb, "10_Ontology_Relations", "subject"):
        if not r.get("subject"):
            continue
        pred_raw, card = r["predicate"], str(r["cardinality"])
        pred = snake(re.sub(r"\(.*?\)", "", pred_raw))
        via = re.search(r"\((.*?)\)", pred_raw)
        ctx = (r["subject"], pred_raw, r["object"])
        targets = resolve(r["object"], ctx)
        right = card.split(":")[1] if ":" in card else card
        for subj in resolve(r["subject"], ctx):
            attrs = classes[subj]["attributes"]
            name = pred if pred not in attrs else f"{pred}_{snake(targets[0])}"
            a = OrderedDict([("description", f"{r['subject']} {pred_raw} {r['object']}")])
            if len(targets) == 1:
                a["range"] = targets[0]
            else:
                a["any_of"] = [{"range": t} for t in targets]
            a["multivalued"] = bool(re.search(r"[NM]", right))
            ann = OrderedDict([("cardinality", card)])
            if via:
                ann["via"] = via.group(1)
            a["annotations"] = ann
            attrs[name] = a
            n_rel += 1

    # Asset carries its identity-map attributes so assets/site01.yaml can be checked with linkml-validate
    asset = classes["Asset"]["attributes"]
    asset.pop("id")
    asset_attrs = OrderedDict([
        ("canonical_asset_id", {"identifier": True, "pattern": r"^AST-[A-Z]{3}-\d{6}$", "description": "Golden asset id AST-<TYPE>-<6 digits>"}),
        ("machine_id", {"required": True, "description": "IIoT machine id (Shopfloor Master Machine_ID)"}),
        ("name", {"required": True}),
        ("process", {"range": "ProcessAreaEnum", "required": True}),
        ("line", {"required": True}),
        ("template", {"range": "AssetTemplateEnum", "required": True}),
        ("mes_asset_id", {"pattern": r"^MES-[A-Z0-9]+-[A-Z0-9_-]+$"}),
        ("cmms_asset_id", {"pattern": r"^EQ-[A-Z0-9]+-[A-Z0-9_-]+$"}),
        ("uns_base_path", {"required": True, "pattern": r"^tyreplant01(/[a-z0-9_]+){3}$"}),
        ("kepware_device", {"required": True}),
        ("criticality", {"range": "CriticalityEnum"}),
        ("match_method", {"range": "MatchMethodEnum"}),
        ("mapping_version", {"required": True}),
        ("mapping_status", {"range": "MappingStatusEnum", "required": True}),
    ])
    classes["Asset"]["attributes"] = OrderedDict(list(asset_attrs.items()) + list(asset.items()))
    classes["AssetRegistry"] = OrderedDict([
        ("description", "Container for one site's asset seed file (assets/<site>.yaml) - validated with linkml-validate"),
        ("tree_root", True),
        ("attributes", OrderedDict([("site", {"required": True}), ("version", {"required": True}),
                                    ("mapping_version", {"required": True}),
                                    ("assets", {"range": "Asset", "multivalued": True, "inlined_as_list": True, "required": True})]))])

    assets = [r for r in rows(wb, "07_Asset_Identity_Map", "Machine_ID (IIoT)") if r.get("Machine_ID (IIoT)")]
    enums = OrderedDict([
        ("ProcessAreaEnum", {"permissible_values": OrderedDict((p, {}) for p in dict.fromkeys(a["process"] for a in assets))}),
        ("AssetTemplateEnum", {"permissible_values": OrderedDict((t, {}) for t in dict.fromkeys(a["template"] for a in assets))}),
        ("CriticalityEnum", {"permissible_values": OrderedDict([("A", {"description": "Critical - stops the plant"}), ("B", {"description": "Important"}), ("C", {"description": "Low impact"})])}),
        ("MatchMethodEnum", {"permissible_values": OrderedDict([("DETERMINISTIC_CROSSWALK", {}), ("ENTITY_RESOLUTION", {}), ("MANUAL", {})])}),
        ("MappingStatusEnum", {"permissible_values": OrderedDict([("PROPOSED", {"description": "Pattern id, not yet confirmed with MES/CMMS"}), ("CONFIRMED", {}), ("INACTIVE", {})])}),
    ])
    schema = OrderedDict([
        ("id", "https://w3id.org/cantier/manufacturing"), ("name", "manufacturing"),
        ("title", "Cantier tyre plant - canonical manufacturing ontology"), ("version", VERSION),
        ("description", "Classes and relations of the tyre plant (sheets 09/10 of workbook v3). Relations are attributes "
                        "on the subject class; cardinality from the workbook is kept as an annotation."),
        ("annotations", {"canonical_model_version": MODEL_VERSION}),
        ("prefixes", OrderedDict([("linkml", "https://w3id.org/linkml/"), ("onto", "https://w3id.org/cantier/onto/")])),
        ("default_prefix", "onto"), ("default_range", "string"), ("imports", ["linkml:types"]),
        ("enums", enums), ("classes", classes)])
    write_yaml("ontology/manufacturing.yaml", schema,
               "# ONTOLOGY (LinkML) - what things exist in the plant and how they connect. Check: python scripts/01_validate_ontology.py\n")
    return classes, assets, n_rel


# ================================================================ B2.2 glossary
SPLIT_ROWS = []
ENUM_TAGS = set()  # measurement types turned into string enums
DP_SHEETS = ["03_IIoT_Datapoints", "04_MES_Datapoints", "05_Inline_Quality", "06_CMMS_Datapoints"]
ENUMS = OrderedDict([
    ("machine_state", ("Operating state of a machine (finding: Boolean / Double could not hold these)", "IIoT / process engineering",
                       OrderedDict([("Stopped", "Not running, no fault"), ("Starting", "Ramping up"), ("Running", "Producing"),
                                    ("Idle", "Ready, not producing"), ("Down", "Fault / breakdown"), ("Alarm", "Running with active alarm")]))),
    ("cure_press_status", ("Curing press state (finding: was a Double in the Tag Dictionary)", "IIoT / process engineering",
                           OrderedDict([("Stopped", "Not running, no fault"), ("Starting", "Ramping up"), ("Running", "Curing"),
                                        ("Idle", "Ready, not curing"), ("Down", "Fault / breakdown"), ("Alarm", "Running with active alarm")]))),
    ("signal_quality", ("OPC UA quality of a reading", "IIoT / process engineering",
                        OrderedDict([("GOOD", "Value can be used"), ("UNCERTAIN", "Use with care"), ("BAD", "Rejected in Silver")]))),
    ("source_system", ("System that produced an event (envelope field)", "Data platform team",
                       OrderedDict([("CANTIER_IIOT", "Cantier IIoT / UNS"), ("MES_EDGE", "MES Edge microservice"), ("MES_CORE", "MES Core DB (DMS CDC)"),
                                    ("CMMS", "Maintenance system"), ("QMS", "Quality system")]))),
    ("tag_domain", ("UNS domain of a tag (5th level of the UNS path)", "IIoT / process engineering",
                    OrderedDict((d, d.replace("_", " ")) for d in ["Context", "Process", "Setpoint", "Condition", "State_Alarm", "Quality",
                                                                    "Utility_Environment", "RawMaterial", "Production"]))),
    ("work_order_type", ("Type of maintenance work order", "Maintenance",
                         OrderedDict([("PREDICTIVE", "Raised from a PdM prediction"), ("PM", "Planned / preventive"), ("CM", "Corrective")]))),
    ("asset_criticality", ("Criticality class of an asset", "Maintenance",
                           OrderedDict([("A", "Critical - stops the plant"), ("B", "Important"), ("C", "Low impact")]))),
    ("cdc_operation", ("AWS DMS change operation", "Data platform team",
                       OrderedDict([("I", "Insert"), ("U", "Update"), ("D", "Delete")]))),
])


def build_glossary(wb, classes):
    measures, entities = OrderedDict(), OrderedDict()
    def all_rows():
        for sheet in DP_SHEETS:
            for r in rows(wb, sheet, "tag_name (source)"):
                yield sheet, r
                while SPLIT_ROWS:
                    yield SPLIT_ROWS.pop(0)

    stop_sheet = None
    for sheet, r in all_rows():
        if True:
            tag = r.get("tag_name (source)")
            if not tag or tag.startswith("—"):
                continue
            if tag.startswith("Part B"):
                stop_sheet = sheet
            if stop_sheet == sheet:
                continue
            gid = r.get("glossary_term_id")
            if not gid:
                continue
            key = gid.split("/")[-1]
            if gid.startswith("onto:entity/"):
                entities.setdefault(key, OrderedDict([
                    ("glossary_term_id", gid), ("label", tag), ("description", f"{tag} - {r.get('silver target') or r.get('CMMS relevance') or r.get('domain')}"),
                    ("data_type", "string"), ("domain", r.get("domain")), ("owner", owner_for(r.get("domain"))), ("status", "active")]))
                continue
            mt_raw = r.get("canonical measurement_type") or ""
            if "/" in mt_raw:  # one workbook row holding two tags, e.g. PLATEN_TEMP_TOP / _BOTTOM
                first, *rest = [x.strip() for x in mt_raw.split("/")]
                base = first.rsplit("_", 1)[0]
                parts = [first] + [base + x if x.startswith("_") else x for x in rest]
                tags = [x.strip() for x in tag.split("/")]
                tbase = tags[0].rsplit("_", 1)[0]
                tags = [tags[0]] + [tbase + x if x.startswith("_") else x for x in tags[1:]]
                for mt, tg in zip(parts, tags):
                    r2 = dict(r, **{"canonical measurement_type": mt, "glossary_term_id": f"onto:measure/{mt.lower()}", "tag_name (source)": tg})
                    SPLIT_ROWS.append((sheet, r2))
                continue
            if key in measures:  # same measurement type in two sheets: keep first, add evidence
                ev = set(measures[key].get("evidence_for", [])) | set(table_refs(r.get("gold usage"), "gold"))
                if ev:
                    measures[key]["evidence_for"] = sorted(ev)
                continue
            dt = dtype(r.get("data_type")) or "string"
            unit, src_unit = r.get("unit (canonical)"), r.get("unit (source)")
            capture = r.get("capture rule") or ""
            issue = r.get("issue / fix")
            status = "derived" if "do not ingest" in capture.lower() else ("missing_at_source" if issue and issue.upper().startswith("MISSING") else "active")
            machine = r.get("source machine type")
            t = OrderedDict([("glossary_term_id", gid), ("label", tag), ("measurement_type", r.get("canonical measurement_type")),
                             ("description", f"{human(tag)} ({r.get('category')}) measured on {machine}" if machine else human(tag)),
                             ("data_type", dt)])
            if unit:
                t["unit"] = unit
            if src_unit and src_unit != unit:
                t["source_unit"] = src_unit
            rng = physical_range(unit, dt)
            if rng:
                t["valid_range"], t["range_basis"] = rng, "physical"
            t["category"] = r.get("category")
            t["domain"] = r.get("domain")
            if capture:
                t["capture_rule"] = capture
            t["observes"] = "Component" if r.get("category") == "Condition" and machine == "Plant Equipment" else "Asset"
            ev = table_refs(r.get("gold usage"), "gold")
            if ev:
                t["evidence_for"] = ev
            tpl = r.get("templates using it")
            if tpl and not tpl.startswith("—"):
                t["templates"] = [x.strip() for x in tpl.split(",")]
            st = table_refs(r.get("silver target"), "silver")
            if st:
                t["silver_tables"] = st
            if r.get("CMMS relevance"):
                t["cmms_relevance"] = r["CMMS relevance"]
            t["status"] = status
            if issue:
                t["source_issue"] = issue
            t["owner"] = owner_for(r.get("domain"))
            measures[key] = t

    # entity terms for every ontology class (id columns in tables point at these)
    for cname, c in classes.items():
        if cname == "AssetRegistry":
            continue
        key = snake(cname)
        entities.setdefault(key, OrderedDict([
            ("glossary_term_id", f"onto:entity/{key}"), ("label", c["annotations"].get("label", cname)),
            ("description", f"Identifier of a {c['annotations'].get('label', cname)}: {c['description']}"),
            ("data_type", "string"), ("ontology_class", cname), ("owner", "Data stewards"), ("status", "active")]))
    for k, v in entities.items():
        v.setdefault("ontology_class", next((c for c in classes if snake(c) == k), None))
        if v["ontology_class"] is None:
            v.pop("ontology_class")

    enums = OrderedDict()
    for k, (desc, owner, vals) in ENUMS.items():
        if k in measures:  # a state tag: make it a string enum (finding fix), keep one term
            m = measures[k]
            m["source_issue"] = f"Source data_type was {m['data_type']} - changed to a string enum. " + (m.get("source_issue") or "")
            m["data_type"], m["enum_values"] = "string", vals
            m.pop("valid_range", None), m.pop("range_basis", None)
            ENUM_TAGS.add(k)
            continue
        enums[k] = OrderedDict([("glossary_term_id", f"onto:enum/{k}"), ("description", desc), ("data_type", "string"),
                                ("enum_values", vals), ("owner", owner), ("status", "active")])
    return measures, entities, enums


# ================================================================ B2.4 tables
LAYER = {"cantier_bronze": "bronze", "cantier_silver": "silver", "cantier_gold": "gold", "cantier_ops": "ops", "cantier_ref": "ref"}
PRIMARY_KEYS = {
    "bronze.iiot_influx_extract": ["bronze_row_id"], "bronze.iiot_uns_events": ["ingestion_id"],
    "bronze.mes_cdc_template": ["dms_lsn"], "bronze.cmms_cdc_template": ["dms_lsn"],
    "silver.iiot_tag_reading": ["canonical_asset_id", "tag_name", "event_ts_utc"], "ops.iiot_tag_reading_rejects": ["bronze_row_id"],
    "silver.iiot_state_event": ["canonical_asset_id", "ts_start"], "silver.mes_production_cycle": ["cycle_id"],
    "silver.qa_inspection_result": ["inspection_id"], "silver.qa_measurement": ["measurement_id"],
    "silver.mdm_asset": ["canonical_asset_id", "valid_from"], "silver.mdm_spec_limit": ["measurement_type", "product_id", "recipe_id", "valid_from"],
    "silver.cmms_work_order": ["work_order_id", "created_at"], "silver.pdm_prediction": ["prediction_id"],
    "silver.mes_lot_action": ["lot_action_id"], "silver.mes_lot": ["lot_id"], "silver.mes_job_order": ["job_order_id"],
    "silver.mes_equipment_problem": ["problem_id"], "silver.qa_ncm": ["ncm_id"],
    "gold.iiot_tag_5min": ["window_start_utc", "canonical_asset_id", "tag_name"],
    "gold.iiot_tag_quality_summary": ["summary_date", "canonical_asset_id", "tag_name"], "gold.maintenance_risk": ["prediction_id"],
    "gold.oee_machine_shift": ["canonical_asset_id", "shift_date", "shift_no"], "gold.curing_cycle_health": ["cycle_id"],
    "gold.process_cycle_summary": ["cycle_id"], "gold.quality_fpy": ["product_id", "canonical_asset_id", "summary_date"],
    "gold.spc_characteristic": ["measurement_type", "canonical_asset_id", "subgroup_id"],
    "gold.downtime_pareto": ["canonical_asset_id", "downtime_reason_code", "summary_date"],
    "gold.asset_reliability": ["canonical_asset_id", "month"], "gold.pdm_feature_hourly": ["canonical_asset_id", "hour_start_utc"],
    "gold.energy_per_tyre": ["canonical_asset_id", "product_id", "summary_date"], "gold.material_consumption_daily": ["material_id", "summary_date"],
    "gold.tyre_traceability": ["tyre_serial"],
}
# grain columns the workbook leaves out of some Gold tables (added so every Gold table has its key)
GOLD_KEY_COLUMNS = {
    "gold.oee_machine_shift": [("canonical_asset_id", "string"), ("shift_date", "date"), ("shift_no", "string")],
    "gold.curing_cycle_health": [("cycle_id", "string"), ("canonical_asset_id", "string")],
    "gold.process_cycle_summary": [("cycle_id", "string"), ("canonical_asset_id", "string"), ("process", "string")],
    "gold.quality_fpy": [("product_id", "string"), ("canonical_asset_id", "string"), ("summary_date", "date")],
    "gold.spc_characteristic": [("measurement_type", "string"), ("canonical_asset_id", "string"), ("subgroup_id", "string")],
    "gold.downtime_pareto": [("canonical_asset_id", "string"), ("downtime_reason_code", "string"), ("summary_date", "date")],
    "gold.asset_reliability": [("canonical_asset_id", "string"), ("month", "date")],
    "gold.pdm_feature_hourly": [("canonical_asset_id", "string"), ("hour_start_utc", "timestamp")],
    "gold.energy_per_tyre": [("canonical_asset_id", "string"), ("product_id", "string"), ("summary_date", "date")],
    "gold.material_consumption_daily": [("material_id", "string"), ("summary_date", "date")],
    "gold.tyre_traceability": [("tyre_serial", "string")],
}
COLUMN_RENAMES = {  # workbook column cells that hold several columns or free text
    ("gold.pdm_feature_hourly", "bearing_temp/vibration stats"): ["bearing_temp_avg", "bearing_temp_max", "bearing_vibration_rms", "bearing_vibration_max"],
    ("gold.pdm_feature_hourly", "gearbox oil temp/vibration"): ["gearbox_oil_temp_avg", "gearbox_oil_temp_max", "gearbox_vibration_rms", "gearbox_vibration_max"],
    ("gold.pdm_feature_hourly", "hydraulic pressure/temp"): ["hydraulic_pressure_avg", "hydraulic_pressure_max", "hydraulic_oil_temp_avg", "hydraulic_oil_temp_max"],
    ("gold.pdm_feature_hourly", "motor temp/current/torque/speed"): ["motor_temperature_avg", "motor_current_avg", "motor_torque_avg", "motor_speed_avg"],
    ("gold.curing_cycle_health", "cure_time_s vs sp"): ["cure_time_s", "cure_time_sp_s"],
}
TECHNICAL = {"bronze_row_id", "event_id", "dms_op", "dms_commit_ts", "dms_lsn", "edge_txn_id", "canonical_model_version",
             "mapping_id", "idempotency_key", "source_time_raw", "source_value_raw", "quality_raw"}
TERM_OVERRIDES = {("silver.cmms_work_order", "work_order_id"): "maintenance_work_order",
                  ("silver.mes_equipment_problem", "work_order_id"): "maintenance_work_order"}
ID_TERMS = {"canonical_asset_id": "asset", "cycle_id": "production_cycle", "cure_cycle_id": "production_cycle",
            "lot_action_id": "production_cycle", "work_order_id": "production_order", "job_order_id": "production_order",
            "lot_id": "lot", "item_id": "product", "product_id": "product", "recipe_id": "recipe", "batch_id": "compound_batch",
            "tyre_serial": "cured_tyre", "shift_no": "shift", "shift_name": "shift", "prediction_id": "prediction",
            "approval_id": "approval", "inspection_id": "inspection_result", "problem_id": "equipment_problem", "ncm_id": "defect",
            "failure_mode": "failure_mode", "component_code": "component", "site_id": "plant", "area": "process_area",
            "line_id": "production_line", "production_line_id": "production_line", "measurement_type": "measurement_type",
            "state": "machine_state", "quality": "signal_quality", "work_order_type": "work_order_type",
            "criticality": "asset_criticality", "operation_id": "operation", "material_id": "material_lot", "subgroup_id": None}


def split_column(table, raw):
    if (table, raw) in COLUMN_RENAMES:
        return COLUMN_RENAMES[(table, raw)]
    if ":" in raw:  # "Banbury: dump_temp, mixing_time, ..."
        return [snake(p) for p in raw.split(":", 1)[1].split(",")]
    if " / " in raw:
        return [snake(p) for p in raw.split(" / ")]
    if "/" in raw:  # "mold_temp_avg/min/max"
        first, *rest = raw.split("/")
        base = first.rsplit("_", 1)[0]
        return [snake(first)] + [snake(f"{base}_{r}") for r in rest]
    return [raw if raw.startswith("_") else snake(raw)]


def infer_type(col):
    c = col.lower()
    if c.startswith("is_") or c.startswith("label_") or c.endswith("_flag") or c.endswith("_pass"):
        return "boolean"
    if c.endswith("_date") or c == "month":
        return "date"
    if c.endswith("_utc") or c.endswith("_ts") or c.startswith("ts_"):
        return "timestamp"
    if c in ("events", "inspected", "passed", "readings", "alarm_count", "nelson_rule_violations") or "count" in c or c.endswith("_cycles") or c.endswith("_at_cure"):
        return "int"
    if c.endswith("_id") or c in ("grade_mix", "top_defect_codes", "failures_by_mode", "rm_lots", "compound_batch", "components",
                                   "green_tyre_id", "build_cycle", "cure_cycle", "press", "inspection_result", "grade", "unit",
                                   "process", "downtime_reason_code", "tyre_serial"):
        return "string"
    return "double"


def infer_unit(col):
    for suf, unit in (("_pct", "%"), ("_s", "s"), ("_min", "min"), ("_h", "h"), ("_hours", "h"), ("_kwh", "kWh"), ("_bar", "bar")):
        if col.endswith(suf):
            return unit
    if col in ("availability", "performance", "quality_rate", "oee", "fpy", "equivalent_cure_pct"):
        return "%"
    return None


def freshness(v):
    m = re.search(r"(\d+)\s*(min|h|d)", v or "")
    return f"{m.group(1)}{m.group(2)}" if m else ({"hourly": "1h", "daily": "1d"}.get((v or "").lower()))


def part_column(p):
    m = re.match(r"\w+\((\w+)\)$", p) or re.match(r"(\w+)$", p)
    return m.group(1) if m else None


def owner_for_table(name):
    short = name.split(".", 1)[1]
    for p, o in (("iiot_", "IIoT / process engineering"), ("mes_", "MES team"), ("qa_", "Quality"), ("cmms_", "Maintenance"),
                 ("mdm_", "Data stewards"), ("pdm_", "Data science")):
        if short.startswith(p):
            return o
    return "Analytics" if name.startswith("gold.") else "Data platform team"
ID_TERMS["process"] = "process_area"


def build_tables(wb, glossary_keys, measures, entities, enums):
    catalogue, specs = OrderedDict(), OrderedDict()
    for r in rows(wb, "12_Glue_Tables", "database"):
        if not r.get("table"):
            continue
        name = f"{LAYER[r['database']]}.{r['table'].replace('<source_table>', 'template')}"
        catalogue[name] = OrderedDict([("grain", r.get("grain")), ("partition", r.get("partition spec (Iceberg hidden)")),
                                       ("source", r.get("source")), ("freshness", r.get("freshness SLA")),
                                       ("glossary_term_id", r.get("glossary_term_id")), ("format", r.get("format"))])
    raw_cols = defaultdict(list)
    for r in rows(wb, "13_Glue_Columns", "database"):
        if r.get("table"):
            raw_cols[f"{LAYER[r['database']]}.{r['table'].replace('<source_table>', 'template')}"].append(r)

    column_terms = OrderedDict()
    all_terms = {**measures, **entities, **enums}
    for name, cols in raw_cols.items():
        layer, short = name.split(".", 1)
        cat = catalogue.get(name, {})
        out, seen = [], set()
        for k, kt in GOLD_KEY_COLUMNS.get(name, []):
            if not any(k in split_column(name, c["column"]) for c in cols):
                cols = [{"column": k, "type": kt, "comment": f"Grain column ({cat.get('grain') or 'key'})"}] + cols
        for c in cols:
            raw = c["column"]
            if raw.startswith("<source"):
                continue
            for col in split_column(name, raw):
                if col in seen:
                    continue
                seen.add(col)
                t = dtype(c.get("type")) if c.get("type") not in (None, "as source") else infer_type(col)
                d = OrderedDict([("name", col), ("type", t), ("nullable", True)])
                gid = c.get("glossary_term_id")
                tech = layer == "bronze" or col.startswith("_") or col in TECHNICAL or (layer == "ops" and col != "reject_reason")
                if tech:
                    d["technical"] = True
                else:
                    term = TERM_OVERRIDES.get((name, col))
                    if not term and gid and not gid.endswith("/*"):
                        term = gid.split("/")[-1]
                    if not term and col in ID_TERMS and ID_TERMS[col]:
                        term = ID_TERMS[col]
                    if not term:  # column term, qualified by table if the same name means another type elsewhere
                        term = col
                        if term in all_terms or (term in column_terms and column_terms[term]["data_type"] != t):
                            if term in all_terms and all_terms[term]["data_type"] == t:
                                pass
                            else:
                                term = f"{short}_{col}"
                        if term not in all_terms and term not in column_terms:
                            ct = OrderedDict([("glossary_term_id", f"onto:column/{term}"),
                                              ("description", c.get("comment") or f"{col.replace('_', ' ').capitalize()} ({cat.get('grain') or short})"),
                                              ("data_type", t)])
                            u = infer_unit(col)
                            if u:
                                ct["unit"] = u
                            ct["owner"] = owner_for_table(name)
                            ct["status"] = "active" if c.get("comment") else "draft"
                            column_terms[term] = ct
                    d["term"] = term
                if c.get("comment") and not tech:
                    d["description"] = c["comment"]
                elif c.get("comment"):
                    d["description"] = c["comment"]
                if c.get("partition col"):
                    d["partition"] = True
                if c.get("pii") == "Y":
                    d["pii"] = True
                out.append(d)
        pk = PRIMARY_KEYS[name]
        for d in out:
            if d["name"] in pk:
                d["nullable"] = False
        parts = [p.strip() for p in re.split(r",\s*(?![^()]*\))", cat.get("partition") or "") if p.strip()]
        names = {d["name"] for d in out}
        parts = [p for p in parts if part_column(p) in names]
        spec = OrderedDict([("name", name), ("layer", layer), ("version", VERSION),
                            ("description", f"{cat.get('grain') or short}. Source: {cat.get('source') or 'see lineage'}"),
                            ("grain", cat.get("grain") or "TBD"), ("primary_key", pk)])
        if parts:
            spec["partition_by"] = parts
        spec["owner"] = owner_for_table(name)
        f = freshness(cat.get("freshness"))
        if f:
            spec["freshness"] = f
        if cat.get("glossary_term_id"):
            spec["glossary_term_id"] = cat["glossary_term_id"]
        spec["source"] = cat.get("source") or "see 17_Lineage"
        spec["status"] = "draft"
        if short.endswith("_template"):
            spec["template"] = True
            spec["notes"] = "One Bronze table per source table (mes_cdc_<source_table>): these DMS columns + the source columns as-is."
        spec["storage"] = {"aws": f"s3://cantier-lake-<acct>/{layer}/{short}/"}
        spec["columns"] = out
        specs[name] = spec
    for name, spec in specs.items():
        write_yaml(f"tables/{spec['layer']}/{name.split('.', 1)[1]}.yaml", spec,
                   "# TABLE SPEC - shape of one table: columns, key, partition, owner. Every column has a glossary term or technical: true.\n")
    cat_out = OrderedDict([("version", VERSION), ("description", "Every table planned in the Glue catalogue (sheet 12). has_spec = a full spec exists in tables/<layer>/."),
                           ("tables", OrderedDict((n, OrderedDict([("has_spec", n in specs)] + [(k, v) for k, v in c.items() if v])) for n, c in catalogue.items()))])
    write_yaml("tables/catalogue.yaml", cat_out, "# TABLE CATALOGUE - all planned tables; specs for the rest arrive with the MES schema export (B2.6) and Gold design (AN).\n")
    return specs, catalogue, column_terms


# ================================================================ B2.3 contracts
ENVELOPE_TYPES = {"string": "string", "timestamp": "timestamp", "enum": "string", "double/string/bool": "any", "int / string": "string", "—": "string"}
ENVELOPE_TERMS = {"source_system": "source_system", "tag_domain": "tag_domain", "quality": "signal_quality", "cycle_id": "production_cycle",
                  "work_order_id": "production_order", "plant_asset_id": "asset", "batch_id": "compound_batch", "recipe_id": "recipe",
                  "product_id": "product", "tyre_serial": "cured_tyre", "shift_no": "shift"}


def build_contracts(wb, specs, templates):
    env_rows = [r for r in rows(wb, "02_Event_Envelope", "field") if r.get("field")]
    fields, example = [], OrderedDict()
    for r in env_rows:
        names = [snake(n) for n in re.split(r"\s*/\s*|,\s*", r["field"].split(":", 1)[-1]) if n.strip()]
        if r["field"].startswith("ingestion"):
            continue  # added at landing (Bronze lineage), not sent by the producer
        req = str(r.get("required") or "N")
        if names[:2] == ["metric_name", "tag_name"]:
            names = ["tag_name", "metric_name"]
        first, aliases = names[0], names[1:]
        if first == "site":  # "site / process / line / machine" are four fields, not aliases
            for n in names:
                fields.append(OrderedDict([("name", n), ("type", "string"), ("required", True), ("description", r.get("definition"))]))
            continue
        if first == "batch_id":
            for n in names:
                f = OrderedDict([("name", n), ("type", "string"), ("required", False), ("description", r.get("definition"))])
                if n in ENVELOPE_TERMS:
                    f["term"] = ENVELOPE_TERMS[n]
                fields.append(f)
            continue
        f = OrderedDict([("name", first), ("type", ENVELOPE_TYPES.get(r.get("type"), r.get("type") or "string")),
                         ("required", req.startswith("Y") and "(" not in req)])
        if "(" in req or "for" in req or "when" in req:
            f["required_when"] = req
        if aliases:
            f["aliases"] = aliases
        if first in ENVELOPE_TERMS:
            f["term"] = ENVELOPE_TERMS[first]
        f["description"] = r.get("definition")
        if r.get("example"):
            f["example"] = str(r["example"])
        fields.append(f)
    envelope = OrderedDict([("id", "envelope.v1"), ("version", VERSION), ("kind", "envelope"),
                            ("description", "Common header of every manufacturing event (UNS -> MSK -> Bronze). Aligned to the PdM reference doc 4.2-4.5."),
                            ("compatibility", "BACKWARD"), ("fields", fields)])
    write_yaml("contracts/envelope.v1.yaml", envelope, "# CONTRACT - the agreed shape of an event between a producer and the platform.\n")

    env_names = {f["name"] for f in fields} | {a for f in fields for a in f.get("aliases", [])}

    def from_table(table, skip=()):
        out = []
        for c in specs[table]["columns"]:
            if c.get("technical") or c["name"] in skip or c["name"] in env_names or c["name"].startswith("_"):
                continue
            f = OrderedDict([("name", c["name"]), ("type", c["type"]), ("required", c["name"] in specs[table]["primary_key"])])
            if c.get("term"):
                f["term"] = c["term"]
            if c.get("pii"):
                f["pii"] = True
            out.append(f)
        return out

    curing = templates["Curing"]
    cond_tags = [t["tag_name"] for t in curing["tags"] if t["domain"] in ("Condition", "Process")]
    pdm_tags = ["SteamValve_Position_CMD", "SteamValve_Position_FB"]
    contracts = [
        OrderedDict([("id", "iiot.reading.v1"), ("version", VERSION), ("kind", "event"), ("extends", "envelope.v1"),
                     ("description", "One tag reading from Cantier IIoT (Process, Setpoint, Condition, Utility, RawMaterial domains)."),
                     ("topic", "manufacturing.<site>.<area>.iiot.reading.v1"), ("key", ["plant_asset_id", "tag_name"]),
                     ("compatibility", "BACKWARD"), ("bronze_table", "bronze.iiot_uns_events"), ("silver_table", "silver.iiot_tag_reading"),
                     ("rules", ["tag_domain in Process, Setpoint, Condition, Utility_Environment, RawMaterial", "unit is the canonical unit of the measurement type"]),
                     ("fields", [])]),
        OrderedDict([("id", "curing.condition.v2"), ("version", VERSION), ("kind", "event"), ("extends", "envelope.v1"),
                     ("description", "Curing press condition/process reading incl. steam valve CMD/FB for the valve-stiction PdM model (PdM doc sample)."),
                     ("topic", "manufacturing.<site>.curing.condition.v2"), ("key", ["plant_asset_id", "tag_name"]),
                     ("compatibility", "BACKWARD"), ("bronze_table", "bronze.iiot_uns_events"), ("silver_table", "silver.iiot_tag_reading"),
                     ("allowed_tags", cond_tags + [t for t in pdm_tags if t not in cond_tags]),
                     ("rules", ["cycle_id required while the press is curing", "SteamValve_Position_CMD/FB are MISSING at source today (finding) - add to Curing template"]),
                     ("fields", [])]),
        OrderedDict([("id", "mes.lot_action.v1"), ("version", VERSION), ("kind", "event"), ("extends", "envelope.v1"),
                     ("description", "One operation executed on a lot (Banbury batch, extrusion run, build, cure, inspection) = production cycle."),
                     ("topic", "manufacturing.<site>.mes.lot_action.v1"), ("key", ["lot_action_id"]), ("compatibility", "BACKWARD"),
                     ("bronze_table", "bronze.mes_edge_events"), ("silver_table", "silver.mes_lot_action"),
                     ("rules", ["lot_action_id is published on UNS as cycle_id at action start (decision 0.1)"]),
                     ("fields", from_table("silver.mes_lot_action", skip=("canonical_asset_id",)))]),
        OrderedDict([("id", "qa.inspection.v1"), ("version", VERSION), ("kind", "event"), ("extends", "envelope.v1"),
                     ("description", "Per-tyre uniformity / balance / dimension / visual result from the uniformity tester or QMS."),
                     ("topic", "manufacturing.<site>.inspection.qa.inspection.v1"), ("key", ["inspection_id"]), ("compatibility", "BACKWARD"),
                     ("bronze_table", "bronze.qa_inspection_events"), ("silver_table", "silver.qa_inspection_result"),
                     ("fields", from_table("silver.qa_inspection_result", skip=("canonical_asset_id",)))]),
    ]
    for c in contracts:
        write_yaml(f"contracts/{c['id']}.yaml", c, "# CONTRACT - the agreed shape of an event between a producer and the platform.\n")
    return [envelope] + contracts


# ================================================================ B2.5 IIoT templates (+ reference instances)
def build_templates(wb, assets, measures):
    inst = [r for r in rows(wb, "08_Tag_Instances", "Machine_ID") if r.get("Machine_ID")]
    by_machine = {a["Machine_ID (IIoT)"]: a for a in assets}
    templates = OrderedDict()
    for r in inst:
        a = by_machine[r["Machine_ID"]]
        tpl = templates.setdefault(a["template"], OrderedDict([("tags", OrderedDict()), ("machines", [])]))
        if r["Machine_ID"] not in tpl["machines"]:
            tpl["machines"].append(r["Machine_ID"])
        if r["tag_name"] not in tpl["tags"]:
            mt = r.get("measurement_type")
            dt = dtype(r["data_type"])
            src_dt = None
            if mt and snake(mt) in ENUM_TAGS:
                src_dt, dt = dt, "string"
            t = OrderedDict([("tag_name", r["tag_name"]), ("measurement_type", snake(mt) if mt else None),
                             ("domain", r["category"]), ("data_type", dt), ("source_data_type", src_dt), ("unit", r.get("unit (canonical)")),
                             ("access", r.get("access")), ("simulation_profile", r.get("simulation profile")),
                             ("bronze_table", (r.get("bronze table") or "").replace("cantier_bronze.", "bronze.")),
                             ("silver_table", r.get("silver table"))])
            templates[a["template"]]["tags"][r["tag_name"]] = OrderedDict((k, v) for k, v in t.items() if v is not None)
    out = OrderedDict()
    for name, tpl in templates.items():
        doc = OrderedDict([("template", name), ("version", VERSION), ("mapping_version", MAPPING_VERSION),
                           ("description", f"{name} machines: which tags they publish and what each tag measures."),
                           ("machines", tpl["machines"]),
                           ("patterns", OrderedDict([
                               ("uns_path", "tyreplant01/{process_lc}/{line_lc}/{machine_lc}/{domain_lc}/{tag_lc}"),
                               ("opc_node_id", "ns=2;s=test.opc.Simulation.Site01.{process}.{line}.{machine}.{domain}.{tag}"),
                               ("kepware_device", "{machine}_Device"), ("kepware_tag", "{tag}")])),
                           ("tags", list(tpl["tags"].values()))])
        out[name] = doc
        write_yaml(f"mappings/iiot/templates/{snake(name)}.yaml", doc,
                   "# IIoT TEMPLATE - tag list of one machine type. machines x tags = tag instances (check: scripts/04_check_coverage.py).\n")
    ref = ROOT / "tests" / "reference" / "tag_instances_v3.csv"
    ref.parent.mkdir(parents=True, exist_ok=True)
    with open(ref, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["machine_id", "canonical_asset_id", "tag_name", "uns_path", "opc_node_id", "kepware_device", "kepware_tag", "measurement_type", "silver_table"])
        for r in inst:
            w.writerow([r["Machine_ID"], r["canonical_asset_id"], r["tag_name"], r["UNS path"], r["opc_node_id"], r["Kepware device"],
                        r["Kepware tag"], r.get("measurement_type"), r.get("silver table")])
    WRITTEN.append("tests/reference/tag_instances_v3.csv")
    return out, len(inst)


# ================================================================ B2.7 assets
def build_assets(assets):
    out = []
    for a in assets:
        out.append(OrderedDict([("canonical_asset_id", a["canonical_asset_id"]), ("machine_id", a["Machine_ID (IIoT)"]),
                                ("name", a["machine name"]), ("process", a["process"]), ("line", a["line"]), ("template", a["template"]),
                                ("mes_asset_id", a["mes_asset_id"]), ("cmms_asset_id", a["cmms_asset_id"]),
                                ("uns_base_path", a["UNS base path"]), ("kepware_device", a["Kepware device"]),
                                ("criticality", a["criticality"]), ("match_method", a["matchMethod"]),
                                ("mapping_version", a["mappingVersion"]), ("mapping_status", a["status"])]))
    doc = OrderedDict([("site", "Site01"), ("version", VERSION), ("mapping_version", assets[0]["mappingVersion"]), ("assets", out)])
    write_yaml("assets/site01.yaml", doc,
               "# ASSET SEEDS - golden id of every machine + its id in MES, CMMS, IIoT, Kepware (-> DynamoDB identity_map, B3.4).\n"
               "# mapping_status PROPOSED = pattern ids from the workbook; replace with real MES/CMMS ids (B2.7, needs decision 0.2).\n"
               "# Check: python scripts/01_validate_ontology.py  (linkml-validate against ontology class AssetRegistry)\n")
    return out


# ================================================================ B2.8 DQ rules (generated from table specs + glossary)
def build_dq(specs, terms):
    count = 0
    for name, spec in specs.items():
        if spec["layer"] not in ("silver", "ops"):
            continue
        cols = {c["name"]: c for c in spec["columns"]}
        short = name.split(".", 1)[1]
        rules = [OrderedDict([("id", f"dq_{short}_pk_unique"), ("type", "unique"), ("columns", spec["primary_key"]), ("severity", "error"),
                              ("description", "One row per grain")])]
        for k in spec["primary_key"]:
            rules.append(OrderedDict([("id", f"dq_{short}_{k}_not_null"), ("type", "not_null"), ("column", k), ("severity", "error")]))
        for c in spec["columns"]:
            t = terms.get(c.get("term") or "")
            if not t:
                continue
            if t.get("enum_values"):
                rules.append(OrderedDict([("id", f"dq_{short}_{c['name']}_valid"), ("type", "accepted_values"), ("column", c["name"]),
                                          ("values", list(t["enum_values"])), ("severity", "warn")]))
            if t.get("valid_range") and c["type"] in ("double", "int", "bigint"):
                r = OrderedDict([("id", f"dq_{short}_{c['name']}_range"), ("type", "range"), ("column", c["name"])])
                r.update(t["valid_range"])
                r["severity"] = "warn"
                rules.append(r)
        if "canonical_asset_id" in cols:
            rules.append(OrderedDict([("id", f"dq_{short}_asset_id_format"), ("type", "pattern"), ("column", "canonical_asset_id"),
                                      ("pattern", r"^AST-[A-Z]{3}-\d{6}$"), ("severity", "error"),
                                      ("description", "Unresolved assets go to ops rejects, never into Silver")]))
        for a, b in (("ts_start", "ts_end"), ("start_ts", "end_ts"), ("planned_start", "planned_end"), ("actual_start", "actual_end"),
                     ("hold_start_ts", "hold_end_ts"), ("reported_ts", "resolved_ts"), ("valid_from", "valid_to"), ("start_ts", "closed_ts")):
            if a in cols and b in cols:
                rules.append(OrderedDict([("id", f"dq_{short}_{a}_before_{b}"), ("type", "expression"), ("check", f"{a} <= {b}"), ("severity", "error")]))
        if "quantity_in" in cols and "quantity_out" in cols:
            rules.append(OrderedDict([("id", f"dq_{short}_out_le_in"), ("type", "expression"), ("check", "quantity_out <= quantity_in"), ("severity", "error")]))
        if {"lsl", "target", "usl"} <= set(cols):
            rules.append(OrderedDict([("id", f"dq_{short}_limits_ordered"), ("type", "expression"), ("check", "lsl <= target and target <= usl"), ("severity", "error")]))
        for c in ("failure_probability_7d",):
            if c in cols:
                rules.append(OrderedDict([("id", f"dq_{short}_{c}_range"), ("type", "range"), ("column", c), ("min", 0), ("max", 1), ("severity", "error")]))
        if short == "iiot_tag_reading":
            rules.append(OrderedDict([("id", "dq_iiot_tag_reading_good_only"), ("type", "accepted_values"), ("column", "quality"), ("values", ["GOOD"]),
                                      ("severity", "error"), ("description", "Silver keeps GOOD readings; others go to ops.iiot_tag_reading_rejects")]))
        seen, uniq = set(), []
        for r in rules:
            if r["id"] not in seen:
                seen.add(r["id"])
                uniq.append(r)
        write_yaml(f"dq/{short}.yaml", OrderedDict([("table", name), ("version", VERSION), ("rules", uniq)]),
                   "# DATA QUALITY - checks each row of this table must pass. error = row rejected, warn = counted. (DQDL generation is B3.6)\n")
        count += 1
    return count


# ================================================================ B2.9 KPIs + B2.6 MES catalogue
def build_mes_catalogue(wb):
    out, in_part_b = [], False
    for r in wb["04_MES_Datapoints"].iter_rows(values_only=True):
        if r[0] == "MES table":
            in_part_b = True
            continue
        if not in_part_b or not r[0]:
            continue
        r = [clean(v) for v in r]
        known = r[2] if r[2] and not r[2].startswith("(confirm") else None
        e = OrderedDict([("mes_tables", [x.strip() for x in r[0].split("/")]), ("module", r[1])])
        if known:
            e["known_columns"] = [x.strip() for x in known.split(",")]
        e.update([("ontology_class", r[3]), ("meaning", r[4]),
                  ("bronze_tables", [x.strip().replace("cantier_bronze.", "") for x in (r[5] or "").split("/") if x.strip()]),
                  ("silver_target", r[6]), ("kpis", [int(x) for x in re.findall(r"\d+", str(r[7] or ""))]), ("relevance", r[8]),
                  ("mapping_status", "waiting_for_schema_export")])
        out.append(OrderedDict((k, v) for k, v in e.items() if v not in (None, [])))
    write_yaml("mappings/mes_catalogue.yaml", OrderedDict([("version", VERSION),
               ("description", "MES DB tables to bring in with DMS CDC (sheet 04 Part B). One mapping file per table goes to mappings/mes/ "
                               "once the MES schema export (INFORMATION_SCHEMA) or sample parquet files are available (B2.6)."),
               ("tables", out)]), "# MES CATALOGUE - which MES tables feed which Silver table and KPIs.\n")
    return out


def build_kpis(wb, mes_catalogue):
    to_silver = {}
    for e in mes_catalogue:
        for t in e["mes_tables"]:
            to_silver[t.lower()] = e.get("silver_target")
    by_area = OrderedDict()
    for r in rows(wb, "04b_MES_KPI_Catalogue", "#"):
        if not isinstance(r.get("#"), int):
            continue
        src_text = r.get("MES source tables") or ""
        mes_tables = [re.sub(r"\(.*?\)", "", x).strip() for x in re.split(r",\s*(?![^()]*\))", src_text) if x.strip()]
        sources = sorted({to_silver[t.lower()] for t in mes_tables if t.lower() in to_silver and to_silver[t.lower()]})
        k = OrderedDict([("id", f"kpi_{snake(r['KPI'])}"), ("number", r["#"]), ("name", r["KPI"]), ("description", r.get("description")),
                         ("formula", r.get("canonical formula")), ("grain", r.get("grain")), ("mes_source_tables", src_text),
                         ("sources", sources), ("gold_table", (r.get("gold product") or "").split(",")[0].strip() or None)])
        if r.get("IIoT / CMMS reconciliation note"):
            k["reconciliation_note"] = r["IIoT / CMMS reconciliation note"]
        k["tyre_relevance"] = r.get("tyre relevance")
        k["owner"] = f"{r['area']} owner (to confirm)"
        k["status"] = "draft"
        by_area.setdefault(r["area"], []).append(OrderedDict((a, b) for a, b in k.items() if b is not None))
    n = 0
    for area, kpis in by_area.items():
        write_yaml(f"kpi/{snake(area)}.yaml", OrderedDict([("area", area), ("version", VERSION), ("kpis", kpis)]),
                   "# KPI DEFINITIONS - formula, grain, sources and owner of each KPI (dbt models generated in B3.7).\n")
        n += len(kpis)
    return n


# ================================================================ main
def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--force", action="store_true", help="overwrite existing definition files")
    a = ap.parse_args()
    if (ROOT / "ontology" / "manufacturing.yaml").exists() and not a.force:
        sys.exit("Definitions already exist - they are maintained in Git now. Use --force only to rebuild drafts from scratch.")
    wb = openpyxl.load_workbook(a.workbook, read_only=True, data_only=True)

    classes, assets, n_rel = build_ontology(wb)
    measures, entities, enums = build_glossary(wb, classes)
    specs, catalogue, column_terms = build_tables(wb, set(), measures, entities, enums)
    all_terms = {**measures, **entities, **enums, **column_terms}
    for fname, terms, desc in (("measurement_types", measures, "What each IIoT / MES / quality / condition datapoint measures (sheets 03-06). "
                                                                 "valid_range = physical bounds only; plant limits live in silver.mdm_spec_limit."),
                               ("entities", entities, "Identifiers of ontology things (asset, cycle, order, lot...). ID columns in tables use these terms."),
                               ("enums", enums, "Closed lists of allowed values."),
                               ("columns", column_terms, "Business columns of Silver / Gold tables that are not a measurement or an entity id.")):
        write_yaml(f"glossary/{fname}.yaml", OrderedDict([("version", VERSION), ("description", desc), ("terms", terms)]),
                   "# GLOSSARY - the dictionary: meaning, type, unit, valid values and owner of each term.\n")
    templates, n_inst = build_templates(wb, assets, measures)
    contracts = build_contracts(wb, specs, templates)
    build_assets(assets)
    n_dq = build_dq(specs, all_terms)
    mes = build_mes_catalogue(wb)
    n_kpi = build_kpis(wb, mes)

    # tables referenced in sheets 03-08 / 04b but missing from the Glue catalogue (sheet 12) -> planned, to design
    referenced = set()
    for folder in ("glossary", "mappings", "contracts", "kpi"):
        for f in (ROOT / folder).rglob("*.yaml"):
            referenced |= set(re.findall(r"\b(?:bronze|silver|gold|ops)\.[a-z][a-z0-9_]*", f.read_text(encoding="utf-8")))
    missing = sorted(referenced - set(catalogue))
    if missing:
        cat_path = ROOT / "tables" / "catalogue.yaml"
        doc = yaml.safe_load(cat_path.read_text(encoding="utf-8"))
        for n in missing:
            doc["tables"][n] = OrderedDict([("has_spec", False), ("note", "Referenced in workbook v3 (sheets 03-08 / 04b) but not in the Glue catalogue (sheet 12) - to design")])
        write_yaml("tables/catalogue.yaml", doc, "# TABLE CATALOGUE - all planned tables; specs for the rest arrive with the MES schema export (B2.6) and Gold design (AN).\n")
    print(f"catalogue  : +{len(missing)} tables referenced in the workbook but missing in sheet 12 (added as planned)")

    findings = [r for r in rows(wb, "01_Findings", "source") if r.get("source")]
    lines = ["# Findings in the source masters (workbook v3, sheet 01)", "",
             "Fix before freezing schemas. Each finding is reflected in the definitions where noted in B2_STATUS.md.", "",
             "| # | Source | Finding | Impact | Fix | Severity |", "|---|---|---|---|---|---|"]
    for i, r in enumerate(findings, 1):
        cells = [str(r.get(k) or "").replace("|", "/").replace("\n", " ") for k in ("source", "finding", "impact", "fix", "severity")]
        lines.append(f"| {i} | " + " | ".join(cells) + " |")
    (ROOT / "docs").mkdir(exist_ok=True)
    (ROOT / "docs" / "FINDINGS.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    WRITTEN.append("docs/FINDINGS.md")

    print(f"ontology   : {len(classes) - 1} classes, {n_rel} relations (+ AssetRegistry container)")
    print(f"glossary   : {len(measures)} measurement types, {len(entities)} entities, {len(enums)} enums, {len(column_terms)} column terms")
    print(f"contracts  : {len(contracts)}")
    print(f"tables     : {len(specs)} specs, {len(catalogue)} in catalogue")
    print(f"templates  : {len(templates)} -> {n_inst} reference tag instances")
    print(f"assets     : {len(assets)}   dq files: {n_dq}   kpis: {n_kpi}   mes tables: {len(mes)}")
    print(f"files      : {len(WRITTEN)}")


if __name__ == "__main__":
    main()
